"""Parse a state education department's staff directory spreadsheet.

Some states (Oklahoma's OSDE, for example) publish official directories of
every school's principal or every district's superintendent, already up to
date, as .xlsx downloads - no scraping needed. This reads either kind of
spreadsheet straight into Contacts, detecting the role (principal,
superintendent, ...) from whichever "<Role>'s Email" column it finds.
"""
import re
from pathlib import Path

from openpyxl import load_workbook

from .emails import clean
from .models import Contact, Target

SCHOOL_HEADER = "school"
DISTRICT_HEADER = "district"
WEBSITE_HEADER = "web site"
COUNTY_HEADER = "county"
CITY_HEADER = "physical city"

_EMAIL_HEADER_RE = re.compile(r"^(.*?)(?:'s)?\s+email$")


def _normalize(value: object) -> str:
    return " ".join(str(value or "").split()).strip().lower()


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    for row in ws.iter_rows(min_row=1, max_row=20):
        headers = {_normalize(cell.value): cell.column - 1 for cell in row if cell.value}
        if any(h.endswith("email") for h in headers):
            return row[0].row, headers
    raise ValueError("No header row with an email column found in the first 20 rows.")


def _column(headers: dict[str, int], contains: str) -> int | None:
    for header, idx in headers.items():
        if contains in header:
            return idx
    return None


def parse_state_directory(path: Path, source_url: str = "") -> list[Contact]:
    """Read a school or district directory .xlsx into one Contact per row with an email."""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    header_row, headers = _find_header_row(ws)
    email_header = next(h for h in headers if h.endswith("email"))
    email_col = headers[email_header]

    role_match = _EMAIL_HEADER_RE.match(email_header)
    role = role_match.group(1).strip() if role_match else ""
    title = role.title()
    name_col = _column(headers, role) if role else None
    school_col = _column(headers, SCHOOL_HEADER)
    district_col = _column(headers, DISTRICT_HEADER)

    contacts = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        email = clean(str(row[email_col])) if row[email_col] else None
        if not email:
            continue
        name = str(row[name_col]).strip() if name_col is not None and row[name_col] else ""
        school = str(row[school_col]).strip() if school_col is not None and row[school_col] else ""
        district = str(row[district_col]).strip() if district_col is not None and row[district_col] else ""
        department = " / ".join(p for p in (school, district) if p)
        contacts.append(Contact(
            email=email,
            name=name,
            title=title,
            department=department,
            source_url=source_url,
        ))
    return contacts


def _normalize_url(url: str) -> str:
    url = url.strip()
    if url and not re.match(r"^https?://", url, re.I):
        url = f"http://{url}"
    return url


def parse_district_targets(path: Path, source: str = "") -> list[Target]:
    """Read an OSDE district directory .xlsx into one Target per district
    that has a website URL - the input to `batch`, not a Contact export."""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    header_row, headers = _find_header_row(ws)
    website_col = _column(headers, WEBSITE_HEADER)
    if website_col is None:
        raise ValueError(f"No {WEBSITE_HEADER!r} column found in {path}.")
    district_col = _column(headers, DISTRICT_HEADER)
    county_col = _column(headers, COUNTY_HEADER)
    city_col = _column(headers, CITY_HEADER)

    targets = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        url = str(row[website_col]).strip() if row[website_col] else ""
        if not url:
            continue
        name = str(row[district_col]).strip() if district_col is not None and row[district_col] else ""
        county = str(row[county_col]).strip() if county_col is not None and row[county_col] else ""
        city = str(row[city_col]).strip() if city_col is not None and row[city_col] else ""
        targets.append(Target(
            name=name,
            homepage_url=_normalize_url(url),
            district=name,
            city=city,
            county=county,
            source=source,
        ))
    return targets
